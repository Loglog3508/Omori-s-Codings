from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


OUTPUT_DIR = Path(r"C:\Users\Administrator\Omori-s-Codings\课程\商务数据分析")
OUTPUT_PATH = OUTPUT_DIR / "大学生哑铃_市场分析数据与计算.xlsx"
QUERY_DATE = "2026-09-10"


TAOBAO_ROWS = [
    [
        1,
        "佑美哑铃男士健身家用锻炼器材可调节重量男杠铃运动套装壶铃宿舍",
        "Umay佑美旗舰店",
        53.75,
        "",
        "本月行业热销",
        0,
        0,
        "平台未展示具体人收货数，不计入保守下限",
    ],
    [
        2,
        "佑美哑铃男士健身家用锻炼器材可调节重量男杠铃健身器材壶铃宿舍",
        "Umay佑美旗舰店",
        88.25,
        "可调节哑铃热卖榜第1名",
        "4000+人收货",
        4000,
        None,
        "以“4000+”的下限4000计",
    ],
    [
        3,
        "裕迅哑铃男士健身家用可调节重量20kg一对男宿舍杠壶铃练组合套装",
        "裕迅旗舰店",
        25.80,
        "可调节健身哑铃热销榜第3名",
        "3000+人收货",
        3000,
        None,
        "以“3000+”的下限3000计",
    ],
    [
        4,
        "朗威哑铃男士健身家用可调节重量器材杠铃力量训练套装壶铃哑铃架",
        "朗威旗舰店",
        89.00,
        "",
        "本月行业热销",
        0,
        0,
        "平台未展示具体人收货数，不计入保守下限",
    ],
    [
        5,
        "哑铃男士健身家用可调节重量10kg20kg一对男学生宿舍杠铃组合套装",
        "超森健身",
        13.80,
        "",
        "2000+人收货",
        2000,
        None,
        "以“2000+”的下限2000计",
    ],
    [
        6,
        "哑铃男士健身家用可调节重量20kg一对男生宿舍杠壶铃锻炼组合套装",
        "裕迅旗舰店",
        18.10,
        "",
        "2000+人收货",
        2000,
        None,
        "以“2000+”的下限2000计",
    ],
    [
        7,
        "哑铃男士健身家用器材可调节重量青少年杠铃壶铃学生宿舍组合套装",
        "多珂星旗舰店",
        18.90,
        "可调节健身哑铃热销榜第7名",
        "1000+人收货",
        1000,
        None,
        "以“1000+”的下限1000计",
    ],
    [
        8,
        "哑铃男士健身器材家用杠铃可调节重量青少年10公斤一对亚哑铃套装",
        "品健运动旗舰店",
        28.80,
        "可调节哑铃热卖榜第2名；行业销量前20",
        "",
        0,
        0,
        "平台未展示具体人收货数，不计入保守下限",
    ],
    [
        9,
        "裕迅哑铃男士健身家用可调节重量20kg一对男宿舍杠铃壶铃组合套装",
        "裕迅旗舰店",
        17.91,
        "入选可调节健身哑铃热销榜",
        "1000+人收货",
        1000,
        None,
        "以“1000+”的下限1000计",
    ],
    [
        10,
        "哑铃男士健身锻炼器材家用可调节重量亚铃男杠铃运动套装壶铃宿舍",
        "品健运动旗舰店",
        49.00,
        "",
        "900+人收货",
        900,
        None,
        "以“900+”的下限900计",
    ],
]


HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
SUBHEADER_FILL = PatternFill("solid", fgColor="D9EAF7")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
NOTE_FILL = PatternFill("solid", fgColor="F2F2F2")
WHITE_FONT = Font(color="FFFFFF", bold=True, name="Arial")
BOLD_FONT = Font(bold=True, name="Arial")
NORMAL_FONT = Font(name="Arial")
INPUT_FONT = Font(name="Arial", color="0000FF")
THIN_GRAY = Side(style="thin", color="B7B7B7")
THIN_BORDER = Border(left=THIN_GRAY, right=THIN_GRAY, top=THIN_GRAY, bottom=THIN_GRAY)


def style_title(ws, cell_range, value):
    ws.merge_cells(cell_range)
    cell = ws[cell_range.split(":")[0]]
    cell.value = value
    cell.font = Font(name="Arial", size=15, bold=True, color="FFFFFF")
    cell.fill = HEADER_FILL
    cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[cell.row].height = 26


def style_table(ws, row, columns):
    for column in range(1, columns + 1):
        cell = ws.cell(row=row, column=column)
        cell.fill = HEADER_FILL
        cell.font = WHITE_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER


def style_cells(ws, min_row, max_row, min_col, max_col):
    for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        for cell in row:
            cell.font = NORMAL_FONT
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = THIN_BORDER


def configure_widths(ws, widths):
    for column, width in widths.items():
        ws.column_dimensions[get_column_letter(column)].width = width


def build_workbook():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    overview = workbook.active
    overview.title = "使用说明"
    taobao = workbook.create_sheet("淘宝TOP10")
    market = workbook.create_sheet("市场规模")
    index = workbook.create_sheet("百度指数录入")
    sources = workbook.create_sheet("数据来源")

    style_title(overview, "A1:F1", "大学生哑铃市场分析：数据、公式与待补采集项")
    overview.append(["文件状态", "查询/制作日期", "核心用途", "已完成数据", "待补数据", "填写说明"])
    style_table(overview, 2, 6)
    overview.append(
        [
            "工作底稿",
            QUERY_DATE,
            "支撑课堂汇报、市场规模计算、红蓝海判断",
            "教育部在校生、国家体育总局产业数据、淘宝销量排序TOP10样本",
            "生意参谋前500、百度指数月度数据、学生问卷购买意愿",
            "黄色底且蓝色文字为需用户填写的输入项；黑色文字为公式结果。",
        ]
    )
    style_cells(overview, 3, 3, 1, 6)
    overview["A5"] = "课程口径"
    overview["A5"].font = BOLD_FONT
    overview["B5"] = "线下市场规模 = 潜在客户数量 × 购买意愿 × 客单价。"
    overview["A6"] = "线上市场口径"
    overview["A6"].font = BOLD_FONT
    overview["B6"] = "淘宝销量排序TOP10样本的可见销量下限 + 生意参谋市场大盘前500品牌成交规模。"
    overview["A7"] = "数据使用警告"
    overview["A7"].font = BOLD_FONT
    overview["B7"] = "淘宝“人收货”不等同于严格的30天销量。所有“+”号均按下限处理；未展示具体量的商品不纳入销量汇总。"
    overview["A8"] = "默认产品边界"
    overview["A8"].font = BOLD_FONT
    overview["B8"] = "暂按“宿舍友好型可调节哑铃套装”分析；产品重量、材质和定价需要小组确认后替换。"
    overview.merge_cells("B5:F5")
    overview.merge_cells("B6:F6")
    overview.merge_cells("B7:F7")
    overview.merge_cells("B8:F8")
    for cell_ref in ("B5", "B6", "B7", "B8"):
        overview[cell_ref].alignment = Alignment(wrap_text=True, vertical="top")
    configure_widths(overview, {1: 18, 2: 20, 3: 28, 4: 34, 5: 32, 6: 34})
    overview.freeze_panes = "A3"

    style_title(taobao, "A1:J1", "淘宝：可调节 哑铃 宿舍 - 销量排序TOP10商品样本")
    taobao["A2"] = "查询关键词"
    taobao["B2"] = "可调节 哑铃 宿舍"
    taobao["D2"] = "查询日期"
    taobao["E2"] = QUERY_DATE
    taobao["G2"] = "排序方式"
    taobao["H2"] = "销量"
    taobao["A3"] = "说明"
    taobao["B3"] = "第1、4、8名仅显示“本月行业热销/行业销量前20”等标签，未显示具体收货数，因此按0纳入下限合计。"
    taobao.merge_cells("B3:J3")
    for cell_ref in ("A2", "D2", "G2", "A3"):
        taobao[cell_ref].font = BOLD_FONT
    for cell_ref in ("B2", "E2", "H2", "B3"):
        taobao[cell_ref].alignment = Alignment(wrap_text=True, vertical="top")
    taobao.append(
        [
            "排序",
            "商品名称",
            "竞品店铺",
            "展示价格（元）",
            "榜单/标签",
            "页面销量信息",
            "销量下限（人）",
            "样本GMV下限（元）",
            "处理说明",
            "数据来源",
        ]
    )
    style_table(taobao, 4, 10)
    for row_index, row_data in enumerate(TAOBAO_ROWS, start=5):
        for column_index, value in enumerate(row_data, start=1):
            taobao.cell(row=row_index, column=column_index, value=value)
        taobao.cell(row=row_index, column=8, value=f"=D{row_index}*G{row_index}")
        taobao.cell(row=row_index, column=10, value=f"淘宝搜索页，关键词“可调节 哑铃 宿舍”，销量排序，{QUERY_DATE}")
    taobao.append(["", "合计/统计", "", "", "", "", "=SUM(G5:G14)", "=SUM(H5:H14)", "", ""])
    summary_row = 15
    for cell in taobao[summary_row]:
        cell.fill = SUBHEADER_FILL
        cell.font = BOLD_FONT
        cell.border = THIN_BORDER
    taobao.append(["", "平均展示价格", "", "=AVERAGE(D5:D14)", "", "", "", "", "", ""])
    taobao.append(["", "中位展示价格", "", "=MEDIAN(D5:D14)", "", "", "", "", "", ""])
    taobao.append(["", "可见销量样本数", "", "", "", "", '=COUNTIF(G5:G14,">0")', "", "", ""])
    style_cells(taobao, 5, 18, 1, 10)
    for row in range(5, 15):
        taobao.cell(row=row, column=4).number_format = '0.00'
        taobao.cell(row=row, column=8).number_format = '#,##0.00'
    taobao["D16"].number_format = '0.00'
    taobao["D17"].number_format = '0.00'
    taobao["G15"].number_format = '#,##0'
    taobao["H15"].number_format = '#,##0.00'
    configure_widths(
        taobao,
        {
            1: 8,
            2: 52,
            3: 20,
            4: 16,
            5: 26,
            6: 18,
            7: 18,
            8: 21,
            9: 34,
            10: 40,
        },
    )
    taobao.freeze_panes = "A5"

    style_title(market, "A1:G1", "大学生哑铃市场规模计算（情景法与平台数据口径）")
    market.append(["模块", "指标/输入", "数值", "单位", "计算或来源", "可否替换", "备注"])
    style_table(market, 2, 7)
    inputs = [
        ["线下市场", "全国普通、职业本专科在校生", 38912600, "人", "教育部《2024年全国教育事业发展统计公报》", "否", "不含研究生、成人教育和网络教育"],
        ["线下市场", "低情景购买意愿", 0.05, "%", "情景假设，需用学生问卷替换", "是", "保守情景"],
        ["线下市场", "基准情景购买意愿", 0.10, "%", "情景假设，需用学生问卷替换", "是", "仅作估算，非调研结论"],
        ["线下市场", "高情景购买意愿", 0.15, "%", "情景假设，需用学生问卷替换", "是", "乐观情景"],
        ["线下市场", "淘宝TOP10平均展示价格", "=淘宝TOP10!D16", "元/件", "淘宝TOP10工作表公式", "是", "展示价格多为低配SKU价格"],
        ["线下市场", "低情景市场规模", "=C3*C4*C7", "元", "潜在客户数 × 购买意愿 × 客单价", "自动", "一人一次购买的年度机会值"],
        ["线下市场", "基准情景市场规模", "=C3*C5*C7", "元", "潜在客户数 × 购买意愿 × 客单价", "自动", "一人一次购买的年度机会值"],
        ["线下市场", "高情景市场规模", "=C3*C6*C7", "元", "潜在客户数 × 购买意愿 × 客单价", "自动", "一人一次购买的年度机会值"],
        ["线上市场", "淘宝TOP10可见销量下限", "=淘宝TOP10!G15", "人/件", "销量排序页可见“人收货”下限", "自动", "不是全行业销量"],
        ["线上市场", "淘宝TOP10样本GMV下限", "=淘宝TOP10!H15", "元", "展示价格 × 可见销量下限", "自动", "不含未显示具体销量的商品"],
        ["线上市场", "生意参谋前500品牌近30天成交额", "", "元", "需用户登录生意参谋市场大盘后录入", "是", "课程要求的行业市场大盘口径"],
        ["线上市场", "行业市场规模（采用生意参谋）", "=C13", "元", "优先采用生意参谋前500汇总值", "自动", "待C13录入；不能用淘宝TOP10样本替代"],
    ]
    for row_data in inputs:
        market.append(row_data)
    style_cells(market, 3, 14, 1, 7)
    for row in range(3, 15):
        market.cell(row=row, column=3).border = THIN_BORDER
    for row in (4, 5, 6, 13):
        market.cell(row=row, column=3).fill = INPUT_FILL
        market.cell(row=row, column=3).font = INPUT_FONT
    for row in (3, 7):
        market.cell(row=row, column=3).font = INPUT_FONT if row == 3 else NORMAL_FONT
    for row in (4, 5, 6):
        market.cell(row=row, column=3).number_format = "0.0%"
    for row in (7, 8, 9, 10, 12, 13, 14):
        market.cell(row=row, column=3).number_format = '#,##0.00'
    market["C3"].comment = Comment("官方统计值：全国普通、职业本专科在校生3891.26万人。", "Codex")
    market["C4"].comment = Comment("请以针对本校/本地大学生的问卷有效样本购买意愿替换。", "Codex")
    market["C5"].comment = Comment("请以针对本校/本地大学生的问卷有效样本购买意愿替换。", "Codex")
    market["C6"].comment = Comment("请以针对本校/本地大学生的问卷有效样本购买意愿替换。", "Codex")
    market["C13"].comment = Comment("生意参谋 > 市场 > 市场大盘，选择相应叶子类目、近30天，累加可见前500品牌成交额后填写。", "Codex")
    market["A16"] = "市场判定（课堂口径）"
    market["A16"].font = BOLD_FONT
    market["B16"] = "电商市场规模10亿元以上为大市场；1000万元至10亿元为中型市场；1000万元以下为小市场。"
    market.merge_cells("B16:G16")
    market["B16"].alignment = Alignment(wrap_text=True)
    configure_widths(market, {1: 14, 2: 34, 3: 20, 4: 14, 5: 36, 6: 12, 7: 38})
    market.freeze_panes = "A3"

    style_title(index, "A1:H1", "百度指数：哑铃 - 月度趋势录入与增长判断")
    index.append(["月份", "哑铃搜索指数", "去年同月指数", "同比增速", "上月指数", "环比增速", "数据来源", "备注"])
    style_table(index, 2, 8)
    for month in range(1, 13):
        row = month + 2
        index.cell(row=row, column=1, value=f"2026-{month:02d}")
        index.cell(row=row, column=4, value=f'=IF(OR(B{row}="",C{row}=""),"",B{row}/C{row}-1)')
        index.cell(row=row, column=6, value=f'=IF(OR(B{row}="",E{row}=""),"",B{row}/E{row}-1)')
        index.cell(row=row, column=7, value="百度指数，关键词“哑铃”")
        for column in (2, 3, 5):
            index.cell(row=row, column=column).fill = INPUT_FILL
            index.cell(row=row, column=column).font = INPUT_FONT
    index["A16"] = "行业生命周期判定规则（按课堂图示）"
    index["A16"].font = BOLD_FONT
    index["B16"] = "增长率 < -10%：衰退期；-10% 至 10%：成熟期；增长率 > 10%：成长期。需要结合至少12个月趋势、同比和环比共同判断。"
    index.merge_cells("B16:H16")
    index["B16"].alignment = Alignment(wrap_text=True)
    style_cells(index, 3, 14, 1, 8)
    for row in range(3, 15):
        index.cell(row=row, column=4).number_format = "0.0%"
        index.cell(row=row, column=6).number_format = "0.0%"
    index["A18"] = "已核验百度指数概览"
    index["A18"].font = BOLD_FONT
    index["A19"] = "查询范围"
    index["B19"] = "整体日均值"
    index["C19"] = "移动日均值"
    index["D19"] = "整体同比"
    index["E19"] = "整体环比"
    index["F19"] = "数据来源"
    index["G19"] = "使用限制"
    for column in range(1, 8):
        index.cell(row=19, column=column).fill = SUBHEADER_FILL
        index.cell(row=19, column=column).font = BOLD_FONT
        index.cell(row=19, column=column).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        index.cell(row=19, column=column).border = THIN_BORDER
    index.append(
        [
            "2026-08-11 ~ 2026-09-09",
            361,
            278,
            -0.29,
            -0.08,
            "百度指数，关键词“哑铃”，全国、PC+移动，查询日2026-09-10",
            "30天短期搜索需求信号，不能单独用于判定全行业衰退。",
            "",
        ]
    )
    index.append(
        [
            "2011-01-01 ~ 2026-09-09（全部）",
            1106,
            716,
            "",
            "",
            "百度指数，关键词“哑铃”，全国、PC+移动，查询日2026-09-10",
            "全周期均值仅作长期背景；生命周期仍需补录近12个月月度序列。",
            "",
        ]
    )
    style_cells(index, 20, 21, 1, 7)
    for row in (20, 21):
        index.cell(row=row, column=4).number_format = "0.0%"
        index.cell(row=row, column=5).number_format = "0.0%"
    configure_widths(index, {1: 13, 2: 16, 3: 18, 4: 14, 5: 16, 6: 14, 7: 28, 8: 38})
    index.freeze_panes = "A3"

    style_title(sources, "A1:F1", "数据来源、取数口径与证据等级")
    sources.append(["编号", "来源", "日期", "可用数据", "证据等级", "使用限制"])
    style_table(sources, 2, 6)
    source_rows = [
        [
            1,
            "教育部《2024年全国教育事业发展统计公报》",
            "2025-06-11",
            "全国普通、职业本专科在校生3891.26万人；普通本科2085.91万人；高职1764.66万人",
            "一级：政府统计",
            "用于潜在客户基数，不代表购买者。",
        ],
        [
            2,
            "国家体育总局、国家统计局《2024年全国体育产业总规模与增加值数据公告》",
            "2025-12-31",
            "体育产业总产出38421亿元；体育用品及相关产品制造增加值4245亿元",
            "一级：政府统计",
            "用于宏观行业背景，不可直接当作哑铃市场规模。",
        ],
        [
            3,
            "淘宝搜索页",
            QUERY_DATE,
            "关键词“可调节 哑铃 宿舍”，销量排序；TOP10商品价格、店铺、销量文本",
            "一级：平台一手页面",
            "“人收货”含“+”号且未明确30天口径；只能作为保守样本下限。",
        ],
        [
            4,
            "生意参谋市场大盘",
            "待补",
            "相应叶子类目近30天、销量/交易额及前500品牌数据",
            "一级：平台一手数据",
            "需登录；本工作簿预留录入单元格。",
        ],
        [
            5,
            "百度指数",
            "2026-09-10查询",
            "关键词“哑铃”，全国、PC+移动、2026-08-11至2026-09-09：整体日均值361、同比-29%、环比-8%",
            "一级：平台一手数据",
            "当前已确认近30天概览；如需年度生命周期曲线，仍需补录至少12个月月度数据。",
        ],
        [
            6,
            "本校/本地大学生问卷",
            "待补",
            "购买意愿、预算、居住条件、重量偏好、品牌偏好",
            "一级：一手调研",
            "建议有效样本不少于30份，并与目标客群对应。",
        ],
        [
            7,
            "国家体育总局：《中国体育用品业年度发展报告（2024）》趋势解读",
            "2025-05-22",
            "2024年全国体育用品电商销售额3337.45亿元；其中健身器材销售额161.06亿元",
            "二级：行业协会年度报告，官方转载",
            "仅作健身器材大类线上市场背景，不能直接替代哑铃或宿舍哑铃的市场规模。",
        ],
        [
            8,
            "国家体育总局：《中国体育用品业年度发展报告（2025）》",
            "2026-05-21",
            "报告称健身器材在2024年高速增长后，2025年小幅回调、呈震荡态势",
            "二级：行业协会年度报告，官方转载",
            "仅用于解释大类趋势分化，哑铃细分仍以淘宝和生意参谋数据为准。",
        ],
    ]
    for row_data in source_rows:
        sources.append(row_data)
    style_cells(sources, 3, 10, 1, 6)
    configure_widths(sources, {1: 8, 2: 42, 3: 16, 4: 52, 5: 20, 6: 42})
    sources.freeze_panes = "A3"

    for ws in workbook.worksheets:
        ws.sheet_view.showGridLines = False
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None and cell.font == Font():
                    cell.font = NORMAL_FONT

    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.save(OUTPUT_PATH)
    print(OUTPUT_PATH)


if __name__ == "__main__":
    build_workbook()
