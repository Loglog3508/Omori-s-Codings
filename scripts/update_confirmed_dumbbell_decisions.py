# -*- coding: utf-8 -*-
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


WORKBOOK = (
    Path(r"C:\Users\Administrator\Omori-s-Codings")
    / "课程"
    / "商务数据分析"
    / "大学生哑铃_市场分析数据与计算.xlsx"
)


def main() -> None:
    workbook = load_workbook(WORKBOOK, data_only=False)
    overview = workbook["使用说明"]
    market = workbook["市场规模"]
    sources = workbook["数据来源"]

    overview["B8"] = (
        "已确认：2-20kg宿舍友好型可调节哑铃套装，静音包胶、可收纳、渐进配重；"
        "主推价59元，价格带49-69元。"
    )

    market["B7"] = "主推客单价（小组确认）"
    market["C7"] = 59
    market["D7"] = "元/套"
    market["E7"] = "小组确认，2026-09-10"
    market["F7"] = "是"
    market["G7"] = "59元为主推价；淘宝TOP10均价40.33元仅作竞品价格参照。"
    market["C7"].fill = PatternFill("solid", fgColor="FFF2CC")
    market["C7"].font = Font(name="Arial", color="0000FF")
    market["C7"].number_format = "0.00"
    market["C7"].comment = Comment(
        "小组于2026-09-10确认主推价59元，价格带49-69元；此值用于线下目标市场的客单价。",
        "Codex",
    )

    expected_formulas = {
        "C8": "=C3*C4*C7",
        "C9": "=C3*C5*C7",
        "C10": "=C3*C6*C7",
    }
    for cell, expected in expected_formulas.items():
        if market[cell].value != expected:
            raise ValueError(f"{cell} formula changed unexpectedly: {market[cell].value!r}")

    sources["C6"] = "2026-09-10"
    sources["D6"] = (
        "已登录有店铺权限账号；市场排行默认类目榜单返回“数据加载错误”，"
        "待恢复后取哑铃叶子类目近30天前500品牌数据"
    )
    sources["F6"] = (
        "账号可进入市场模块，但平台榜单暂未正常返回数据；"
        "不可用淘宝TOP10样本替代前500行业大盘。"
    )

    if "生意参谋TOP10（7天）" in workbook.sheetnames:
        del workbook["生意参谋TOP10（7天）"]
    sycm = workbook.create_sheet("生意参谋TOP10（7天）", 2)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    subheader_fill = PatternFill("solid", fgColor="D9EAF7")
    white_font = Font(name="Arial", color="FFFFFF", bold=True)
    normal_font = Font(name="Arial")
    bold_font = Font(name="Arial", bold=True)
    thin = Side(style="thin", color="B7B7B7")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    sycm.merge_cells("A1:I1")
    sycm["A1"] = "生意参谋：哑铃类目交易总量TOP10（7天可见样本）"
    sycm["A1"].fill = header_fill
    sycm["A1"].font = Font(name="Arial", size=15, bold=True, color="FFFFFF")
    sycm["A1"].alignment = Alignment(horizontal="center", vertical="center")
    sycm.row_dimensions[1].height = 26

    sycm.append(["类目路径", "统计周期", "榜单口径", "可见页数", "", "", "", "", ""])
    sycm.merge_cells("D2:I2")
    sycm["A2"] = "运动/瑜伽/健身/球迷用品 > 健身训练器材 > 哑铃"
    sycm["B2"] = "2026-09-03 至 2026-09-09"
    sycm["C2"] = "交易总量"
    sycm["D2"] = "每页10条、页面显示30页；至少可见TOP300商品。当前仅记录第一页TOP10。"
    for cell in ("A2", "B2", "C2", "D2"):
        sycm[cell].font = bold_font if cell != "D2" else normal_font
        sycm[cell].alignment = Alignment(wrap_text=True, vertical="top")

    sycm.append(["重要限制", "支付买家数为区间；不等于销量，也不是成交额。当前为7天商品TOP10，不能替代近30天行业前500品牌成交额。", "", "", "", "", "", "", ""])
    sycm.merge_cells("B3:I3")
    sycm["A3"].font = bold_font
    sycm["B3"].alignment = Alignment(wrap_text=True, vertical="top")

    headers = [
        "排名",
        "商品",
        "商品关键词",
        "店铺",
        "支付买家数下限",
        "支付买家数上限",
        "访客数下限",
        "访客数上限",
        "排行变化",
    ]
    sycm.append(headers)
    for cell in sycm[4]:
        cell.fill = header_fill
        cell.font = white_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    top10 = [
        [1, "Keep哑铃女士健身家用男士力量训练器材锻炼初学者练臂肌亚铃一对", "女士 / 健身 / 家用", "Keep官方旗舰店", 5000, 7500, 100000, 250000, "持平"],
        [2, "佑美哑铃男士健身家用锻炼器材可调节重量男杠铃健身器材壶铃宿舍", "健身 / 家用 / 男士", "Umay佑美旗舰店", 1000, 2500, 25000, 50000, "持平"],
        [3, "朗威哑铃女士健身家用器材儿童小重量2kg一对男士青少年组合套装", "女士 / 健身 / 家用", "朗威旗舰店", 2500, 5000, 25000, 50000, "持平"],
        [4, "哑铃健身男士家用六角5kg一对6公斤器材锻炼包胶亚铃女士青少年", "健身 / 家用 / 男士", "艾美仕运动旗舰店", 2500, 5000, 25000, 50000, "升2名"],
        [5, "哑铃男士健身器材家用杠铃可调节重量青少年10公斤一对亚哑铃套装", "健身 / 家用 / 男士", "品健运动旗舰店", 2500, 5000, 25000, 50000, "持平"],
        [6, "Keep软壶铃女士健身家用课程同款哑铃男士运动塑形深蹲提壶器材", "健身 / 家用 / 女士", "Keep官方旗舰店", 1000, 2500, 10000, 25000, "降2名"],
        [7, "【芭芭农场】1kg哑铃一对", "农场 / 一对 / 瑜伽", "潮达商贸企业店", 10000, 25000, 25000, 50000, "升2名"],
        [8, "Keep哑铃健身男女家用锻炼运动器械哑铃实心铸铁宿舍亚铃10kg一对", "健身 / 男士 / 家用", "Keep官方旗舰店", 750, 1000, 10000, 25000, "升2名"],
        [9, "迪卡侬哑铃女士健身家用男士一对居家健身器材训练铸铁小哑铃RFJ3", "女士 / 健身 / 家用", "迪卡侬旗舰店", 1000, 2500, 25000, 50000, "降2名"],
        [10, "裕迅哑铃男士健身家用可调节重量20kg一对男宿舍杠壶铃练组合套装", "健身 / 男士 / 家用", "裕迅旗舰店", 2500, 5000, 10000, 25000, "升3名"],
    ]
    for row in top10:
        sycm.append(row)
    sycm.append(["合计", "", "", "", "=SUM(E5:E14)", "=SUM(F5:F14)", "=SUM(G5:G14)", "=SUM(H5:H14)", ""])
    for cell in sycm[15]:
        cell.fill = subheader_fill
        cell.font = bold_font
        cell.border = border

    for row in sycm.iter_rows(min_row=5, max_row=15, min_col=1, max_col=9):
        for cell in row:
            cell.font = normal_font if cell.row != 15 else bold_font
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border
    for col in range(5, 9):
        for row in range(5, 16):
            sycm.cell(row=row, column=col).number_format = "#,##0"
    for col, width in {
        1: 9,
        2: 52,
        3: 22,
        4: 22,
        5: 18,
        6: 18,
        7: 16,
        8: 16,
        9: 12,
    }.items():
        sycm.column_dimensions[get_column_letter(col)].width = width
    sycm.freeze_panes = "A5"
    sycm.sheet_view.showGridLines = False

    sources.append(
        [
            9,
            "生意参谋市场排行",
            "2026-09-10查询",
            "运动/瑜伽/健身/球迷用品 > 健身训练器材 > 哑铃；2026-09-03至2026-09-09交易总量TOP10商品，支付买家数和访客数区间",
            "一线：平台一手数据",
            "当前为7天商品TOP10可见样本；支付买家数为区间，不可替代近30天前500品牌成交额。",
        ]
    )

    if "生意参谋店铺TOP10（30天）" in workbook.sheetnames:
        del workbook["生意参谋店铺TOP10（30天）"]
    shops = workbook.create_sheet("生意参谋店铺TOP10（30天）", 3)
    shops.merge_cells("A1:G1")
    shops["A1"] = "生意参谋：哑铃类目店铺交易总量TOP10（近30天）"
    shops["A1"].fill = header_fill
    shops["A1"].font = Font(name="Arial", size=15, bold=True, color="FFFFFF")
    shops["A1"].alignment = Alignment(horizontal="center", vertical="center")
    shops.row_dimensions[1].height = 26
    shops.append(["类目路径", "统计周期", "榜单口径", "可见页数", "", "", ""])
    shops.merge_cells("D2:G2")
    shops["A2"] = "运动/瑜伽/健身/球迷用品 > 健身训练器材 > 哑铃"
    shops["B2"] = "2026-08-11 至 2026-09-09"
    shops["C2"] = "交易总量（店铺）"
    shops["D2"] = "每页10条、页面显示30页；至少可见TOP300店铺。当前仅记录第一页TOP10。"
    shops.append(["重要限制", "支付金额仅显示排名、未显示数值；不可计算前500品牌成交额。访客数为区间。", "", "", "", "", ""])
    shops.merge_cells("B3:G3")
    for cell in ("A2", "B2", "C2", "A3"):
        shops[cell].font = bold_font
    for cell in ("A2", "B2", "C2", "D2", "B3"):
        shops[cell].alignment = Alignment(wrap_text=True, vertical="top")
    shop_headers = ["排名", "店铺", "支付金额排名", "访客数下限", "访客数上限", "在售商品数量", "排行变化"]
    shops.append(shop_headers)
    for cell in shops[4]:
        cell.fill = header_fill
        cell.font = white_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    shop_rows = [
        [1, "Keep官方旗舰店", 1, 500000, 750000, 10, "持平"],
        [2, "Umay佑美旗舰店", 2, 250000, 500000, 20, "持平"],
        [3, "朗威旗舰店", 3, 750000, 1000000, 120, "持平"],
        [4, "proiron运动旗舰店", 4, 100000, 250000, 30, "升1名"],
        [5, "cardlosjohan旗舰店", 5, 250000, 500000, 80, "降1名"],
        [6, "品健运动旗舰店", 6, 100000, 250000, 70, "升1名"],
        [7, "YOTTOY官方旗舰店", 7, 100000, 250000, 30, "降1名"],
        [8, "迪卡侬旗舰店", 8, 100000, 250000, 10, "升3名"],
        [9, "艾美仕运动旗舰店", 9, 100000, 250000, 60, "升7名"],
        [10, "裕迅旗舰店", 10, 75000, 100000, 30, "持平"],
    ]
    for row in shop_rows:
        shops.append(row)
    shops.append(["合计", "", "", "=SUM(D5:D14)", "=SUM(E5:E14)", "=SUM(F5:F14)", ""])
    for row in shops.iter_rows(min_row=5, max_row=15, min_col=1, max_col=7):
        for cell in row:
            cell.font = bold_font if cell.row == 15 else normal_font
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border
            if cell.column in (3, 4, 5, 6):
                cell.number_format = "#,##0"
    for cell in shops[15]:
        cell.fill = subheader_fill
    for col, width in {1: 9, 2: 28, 3: 16, 4: 16, 5: 16, 6: 18, 7: 12}.items():
        shops.column_dimensions[get_column_letter(col)].width = width
    shops.freeze_panes = "A5"
    shops.sheet_view.showGridLines = False
    sources.append(
        [
            10,
            "生意参谋市场排行",
            "2026-09-10查询",
            "哑铃类目2026-08-11至2026-09-09交易总量TOP10店铺：支付金额排名、访客数区间、在售商品数量",
            "一线：平台一手数据",
            "仅有支付金额排名，没有金额值；可用于头部店铺与流量集中度分析，不能替代前500品牌成交额。",
        ]
    )

    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.save(WORKBOOK)


if __name__ == "__main__":
    main()
