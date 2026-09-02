# -*- coding: utf-8 -*-
"""导出采集到的天气数据为 Excel 报表，输出到 report_charts 文件夹。"""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
RAW_FILE = ROOT / "data" / "weather_raw.json"
DATA_FILE = ROOT / "data" / "weather_clean.json"
META_FILE = ROOT / "data" / "weather_metadata.json"
OUT_DIR = ROOT / "report_charts"
EXCEL_FILE = OUT_DIR / "weather_report.xlsx"


HEADER_FONT = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
UNIT_FONT = Font(name="微软雅黑", size=9, color="999999")
TITLE_FONT = Font(name="微软雅黑", size=16, bold=True, color="1F2933")
SUB_TITLE_FONT = Font(name="微软雅黑", size=11, bold=True, color="2F5496")
BODY_FONT = Font(name="微软雅黑", size=10)
THIN_BORDER = Border(
    left=Side(style="thin", color="D0D0D0"),
    right=Side(style="thin", color="D0D0D0"),
    top=Side(style="thin", color="D0D0D0"),
    bottom=Side(style="thin", color="D0D0D0"),
)
CENTER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT_ALIGN = Alignment(horizontal="left", vertical="center")

RAW_FIELD_CN = {
    "time": "日期",
    "weather_code": "天气代码",
    "temperature_2m_max": "最高气温",
    "temperature_2m_min": "最低气温",
    "temperature_2m_mean": "平均气温",
    "apparent_temperature_max": "体感最高温",
    "apparent_temperature_min": "体感最低温",
    "precipitation_sum": "降水量",
    "rain_sum": "降雨量",
    "snowfall_sum": "降雪量",
    "precipitation_hours": "降水时长",
    "wind_speed_10m_max": "最大风速",
    "wind_gusts_10m_max": "最大阵风",
    "shortwave_radiation_sum": "太阳辐射总量",
}


def style_header(ws, row, col_count):
    for col in range(1, col_count + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER_ALIGN
        cell.border = THIN_BORDER


def style_body(ws, row, col_count, alignment=CENTER_ALIGN):
    for col in range(1, col_count + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = BODY_FONT
        cell.alignment = alignment
        cell.border = THIN_BORDER


def auto_width(ws, col_count, min_width=12, max_width=36):
    for col in range(1, col_count + 1):
        letter = get_column_letter(col)
        best = min_width
        for row in ws.iter_rows(min_col=col, max_col=col, values_only=True):
            for cell in row:
                if cell is not None:
                    text = str(cell)
                    width = sum(2 if ord(c) > 127 else 1 for c in text)
                    best = max(best, width + 4)
        ws.column_dimensions[letter].width = min(best, max_width)


def build_raw_data_sheet(wb):
    ws = wb.create_sheet("爬取原始数据")

    raw = json.loads(RAW_FILE.read_text(encoding="utf-8-sig"))
    daily = raw["daily"]
    units = raw.get("daily_units", {})
    fields = list(daily.keys())

    for col, field in enumerate(fields, 1):
        ws.cell(row=1, column=col, value=RAW_FIELD_CN.get(field, field))
    style_header(ws, 1, len(fields))

    for col, field in enumerate(fields, 1):
        unit = units.get(field, "")
        label = f"{RAW_FIELD_CN.get(field, field)}" if not unit else f"{RAW_FIELD_CN.get(field, field)}（{unit}）"
        cell = ws.cell(row=2, column=col, value=label)
        cell.font = UNIT_FONT
        cell.alignment = CENTER_ALIGN
        cell.border = THIN_BORDER

    total = len(daily[fields[0]])
    for col, field in enumerate(fields, 1):
        for i, value in enumerate(daily[field]):
            ws.cell(row=i + 3, column=col, value=value if value is not None else "")

    for row in range(3, total + 3):
        style_body(ws, row, len(fields))

    auto_width(ws, len(fields))
    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:{get_column_letter(len(fields))}{total + 2}"


def build_summary_sheet(wb, meta):
    ws = wb.active
    ws.title = "汇总概览"

    ws.merge_cells("A1:C1")
    cell = ws["A1"]
    cell.value = "大连市近10年天气数据分析报告"
    cell.font = TITLE_FONT
    cell.alignment = CENTER_ALIGN

    ws.merge_cells("A2:C2")
    ws["A2"].value = f"生成时间：{meta.get('generated_at', '')}"
    ws["A2"].font = Font(name="微软雅黑", size=9, color="808080")
    ws["A2"].alignment = CENTER_ALIGN

    headers = ["指标", "数值", "说明"]
    for col, h in enumerate(headers, 1):
        ws.cell(row=4, column=col, value=h)
    style_header(ws, 4, 3)

    rows_data = [
        ("数据来源", meta.get("source", ""), "Open-Meteo Archive API"),
        ("数据范围", f"{meta.get('date_start', '')} ~ {meta.get('date_end', '')}", ""),
        ("总天数", f"{meta.get('total_days', 0)} 天", ""),
        ("平均气温", f"{meta.get('avg_temp', 0):.2f} °C", "近10年日均气温的均值"),
        ("最高气温", f"{meta.get('max_temp', 0):.2f} °C", "记录中的单日最高气温"),
        ("最低气温", f"{meta.get('min_temp', 0):.2f} °C", "记录中的单日最低气温"),
        ("累计降水量", f"{meta.get('total_precipitation', 0):.2f} mm", "近10年降水总量"),
        ("降水天数", f"{meta.get('rainy_days', 0)} 天", "降水量 ≥ 0.1mm 的天数"),
        ("高温天数（≥30°C）", f"{meta.get('hot_days', 0)} 天", ""),
        ("低温天数（≤-5°C）", f"{meta.get('cold_days', 0)} 天", ""),
    ]

    for i, (indicator, value, note) in enumerate(rows_data, 5):
        ws.cell(row=i, column=1, value=indicator)
        ws.cell(row=i, column=2, value=value)
        ws.cell(row=i, column=3, value=note)
        style_body(ws, i, 3, LEFT_ALIGN)
        ws.cell(row=i, column=1).font = SUB_TITLE_FONT

    auto_width(ws, 3)


def build_daily_sheet(wb, rows):
    ws = wb.create_sheet("每日数据（清洗后）")

    headers = [
        "日期", "年份", "月份", "季节",
        "天气代码", "天气描述",
        "最高温（°C）", "最低温（°C）", "平均温（°C）", "温差（°C）",
        "体感最高温（°C）", "体感最低温（°C）",
        "降水量（mm）", "降雨量（mm）", "降雪量（mm）", "降水时长（h）",
        "最大风速（km/h）", "最大阵风（km/h）", "太阳辐射（MJ/m²）",
        "高温日", "低温日", "降雨日", "大雨日", "降雪日", "大风日",
    ]

    for col, h in enumerate(headers, 1):
        ws.cell(row=1, column=col, value=h)
    style_header(ws, 1, len(headers))

    for i, row in enumerate(rows, 2):
        values = [
            row.get("date", ""), row.get("year", ""), row.get("month", ""), row.get("season", ""),
            row.get("weather_code", ""), row.get("weather_text", ""),
            row.get("temp_max", ""), row.get("temp_min", ""), row.get("temp_mean", ""), row.get("temp_range", ""),
            row.get("apparent_max", ""), row.get("apparent_min", ""),
            row.get("precipitation", ""), row.get("rain", ""), row.get("snowfall", ""), row.get("precipitation_hours", ""),
            row.get("wind_max", ""), row.get("wind_gust", ""), row.get("radiation", ""),
            "是" if row.get("is_hot") else "否",
            "是" if row.get("is_cold") else "否",
            "是" if row.get("is_rainy") else "否",
            "是" if row.get("is_heavy_rain") else "否",
            "是" if row.get("is_snowy") else "否",
            "是" if row.get("is_windy") else "否",
        ]
        for col, val in enumerate(values, 1):
            ws.cell(row=i, column=col, value=val)
        style_body(ws, i, len(headers))

    auto_width(ws, len(headers))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"


def build_yearly_sheet(wb, meta):
    ws = wb.create_sheet("年度统计")

    headers = ["年份", "平均气温（°C）", "降水量（mm）", "高温天数", "低温天数", "降雨天数"]
    for col, h in enumerate(headers, 1):
        ws.cell(row=1, column=col, value=h)
    style_header(ws, 1, len(headers))

    for i, item in enumerate(meta.get("yearly", []), 2):
        values = [
            item.get("year", ""), item.get("avg_temp", ""), item.get("precipitation", ""),
            item.get("hot_days", ""), item.get("cold_days", ""), item.get("rainy_days", ""),
        ]
        for col, val in enumerate(values, 1):
            ws.cell(row=i, column=col, value=val)
        style_body(ws, i, len(headers))

    auto_width(ws, len(headers))
    ws.freeze_panes = "A2"


def build_monthly_sheet(wb, meta):
    ws = wb.create_sheet("月度统计")

    headers = ["月份", "平均气温（°C）", "降水量（mm）", "天数"]
    for col, h in enumerate(headers, 1):
        ws.cell(row=1, column=col, value=h)
    style_header(ws, 1, len(headers))

    for i, item in enumerate(meta.get("monthly", []), 2):
        values = [item.get("month", ""), item.get("avg_temp", ""), item.get("precipitation", ""), item.get("days", "")]
        for col, val in enumerate(values, 1):
            ws.cell(row=i, column=col, value=val)
        style_body(ws, i, len(headers))

    auto_width(ws, len(headers))
    ws.freeze_panes = "A2"


def build_weather_type_sheet(wb, meta):
    ws = wb.create_sheet("天气类型统计")

    headers = ["天气类型", "出现天数", "占比（%）"]
    for col, h in enumerate(headers, 1):
        ws.cell(row=1, column=col, value=h)
    style_header(ws, 1, len(headers))

    total = sum(item["value"] for item in meta.get("weather_counts", []))
    for i, item in enumerate(meta.get("weather_counts", []), 2):
        pct = round(item["value"] / total * 100, 2) if total else 0
        values = [item.get("name", ""), item.get("value", 0), pct]
        for col, val in enumerate(values, 1):
            ws.cell(row=i, column=col, value=val)
        style_body(ws, i, len(headers))
        ws.cell(row=i, column=3).number_format = "0.00"

    auto_width(ws, len(headers))
    ws.freeze_panes = "A2"


def main():
    rows = json.loads(DATA_FILE.read_text(encoding="utf-8-sig"))
    meta = json.loads(META_FILE.read_text(encoding="utf-8-sig"))

    OUT_DIR.mkdir(parents=True, exist_ok=True)

    wb = Workbook()

    build_raw_data_sheet(wb)
    build_summary_sheet(wb, meta)
    build_daily_sheet(wb, rows)
    build_yearly_sheet(wb, meta)
    build_monthly_sheet(wb, meta)
    build_weather_type_sheet(wb, meta)

    wb.save(EXCEL_FILE)
    print(f"Excel report saved to {EXCEL_FILE}")


if __name__ == "__main__":
    main()